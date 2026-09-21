"""Own Product Document Integration 无网络 Controlled E2E。"""

from __future__ import annotations

import json
import shutil
import sys
import tempfile
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from knowledge import KnowledgeRetriever, KnowledgeStore
from llm.retrieval_models import EmbeddingOutput, RerankItem, RerankOutput
from pipeline.documents import (
    OwnProductDocumentService, product_scope_id, workspace_knowledge_path,
)


class ControlledEmbedding:
    space_id = "controlled/own-product-document-e2e-v1"

    def __init__(self):
        self.calls = 0

    def embed(self, texts):
        self.calls += 1
        return EmbeddingOutput([[1.0, 0.0] for _ in texts], [])


class ControlledRerank:
    def rerank(self, query, documents, top_k):
        return RerankOutput(
            [RerankItem(index=index, score=1.0 - index / 1000)
             for index in range(min(top_k, len(documents)))], []
        )


def make_docx(path: Path, sections: list[tuple[str, str]]) -> Path:
    """生成受控 DOCX，只用于证明结构缩减时旧 Section 不会残留。"""
    document = Document()
    for heading, body in sections:
        document.add_heading(heading, level=1)
        document.add_paragraph(body)
    document.save(path)
    return path


def main() -> int:
    report = {}
    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary)
        state = root / "state"
        workspace, product = "controlled-workspace", "controlled-product"
        scope = product_scope_id(workspace, product)
        embedding = ControlledEmbedding()
        service = OwnProductDocumentService(state, embedding_factory=lambda: embedding)
        fixture = ROOT / "experiments" / "fixtures" / "file_parsing" / "demo.docx"

        added = service.add_own_product_document(
            workspace_id=workspace, product_id=product, file=fixture
        )
        store = KnowledgeStore(workspace_knowledge_path(state, workspace))
        retriever = KnowledgeRetriever(store, embedding, ControlledRerank())
        retrieved = retriever.retrieve(
            "竞争事件证据追踪", knowledge_domain="own_product", top_n=20, top_k=20,
            filters={"product_id": scope},
        )
        fixture_results = [item for item in retrieved.results
                           if item.document_id.startswith(added.document_id + ":unit:")]
        report["case_a"] = {
            "add_completed": added.ingestion_status == "completed",
            "retrieved": bool(fixture_results),
            "reranked": bool(fixture_results and fixture_results[0].rerank_score is not None),
            "provenance_preserved": bool(
                fixture_results and fixture_results[0].metadata.get("heading_path")
                and fixture_results[0].metadata.get("source_file_sha256")
            ),
        }

        old = make_docx(root / "old.docx", [
            ("一", "OLD_ONE"), ("二", "OLD_TWO"), ("三", "STALE_SECTION")
        ])
        new = make_docx(root / "new.docx", [("一", "NEW_ONE"), ("二", "CURRENT_SECTION")])
        structural = service.add_own_product_document(
            workspace_id=workspace, product_id=product, file=old,
            display_name="structural.docx",
        )
        replaced = service.replace_own_product_document(
            workspace_id=workspace, product_id=product,
            document_id=structural.document_id, new_file=new,
        )
        after = retriever.retrieve(
            "当前结构", knowledge_domain="own_product", top_n=50, top_k=50,
            filters={"product_id": scope},
        )
        contents = "\n".join(item.content for item in after.candidates)
        report["case_b"] = {
            "same_document_id": replaced.document_id == structural.document_id,
            "three_to_two": structural.unit_count == 3 and replaced.unit_count == 2,
            "stale_unit_deleted": store.document_record(
                "own_product", structural.document_id + ":unit:3"
            ) is None,
            "old_content_absent": "STALE_SECTION" not in contents,
            "new_content_available": "CURRENT_SECTION" in contents,
        }

        copied = root / "copied-demo.docx"
        shutil.copyfile(fixture, copied)
        before_counts, before_calls = store.counts(), embedding.calls
        duplicate = service.add_own_product_document(
            workspace_id=workspace, product_id=product, file=copied
        )
        report["case_c"] = {
            "existing_document_reused": duplicate.document_id == added.document_id,
            "existing_status": duplicate.content_status == "existing",
            "no_duplicate_knowledge": store.counts() == before_counts,
            "no_duplicate_embedding": embedding.calls == before_calls,
        }

        reopened = OwnProductDocumentService(state, embedding_factory=lambda: embedding)
        reopened_status = reopened.get_own_product_document(
            workspace_id=workspace, product_id=product, document_id=structural.document_id
        )
        reopened_store = KnowledgeStore(workspace_knowledge_path(state, workspace))
        reopened_results = KnowledgeRetriever(
            reopened_store, embedding, ControlledRerank()
        ).retrieve(
            "CURRENT_SECTION", knowledge_domain="own_product", top_n=50, top_k=50,
            filters={"product_id": scope},
        )
        report["case_d"] = {
            "document_id_stable": reopened_status.document_id == structural.document_id,
            "current_state_persisted": reopened_status.unit_count == 2,
            "retrieval_persisted": any(
                "CURRENT_SECTION" in item.content for item in reopened_results.results
            ),
        }

    report["status"] = "PASS" if all(
        all(case.values()) for name, case in report.items() if name.startswith("case_")
    ) else "FAIL"
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
