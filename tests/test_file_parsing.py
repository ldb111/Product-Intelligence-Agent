"""File Parsing 控制测试：Fixture 全为模拟资料，不读取真实用户文件。"""

import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from file_parsing import parse_and_prepare, parse_document, prepare_for_knowledge
from knowledge.chunking import chunk_document


FIXTURES = Path(__file__).resolve().parents[1] / "experiments" / "fixtures" / "file_parsing"


class FileParsingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def write(self, name, data, *, binary=False):
        path = self.root / name
        (path.write_bytes(data) if binary else path.write_text(data, encoding="utf-8"))
        return path

    def docx(self, name="sample.docx", build=None):
        path = self.root / name
        document = Document()
        if build:
            build(document)
        document.save(path)
        return path

    def xlsx(self, name="sample.xlsx", build=None):
        path = self.root / name
        workbook = Workbook()
        if build:
            build(workbook)
        workbook.save(path)
        return path

    def test_txt_normal(self):
        result = parse_document(self.write("a.txt", "产品支持证据追踪。"))
        self.assertEqual(result.status, "parsed")
        self.assertEqual(result.document.units[0].text, "产品支持证据追踪。")

    def test_txt_empty(self):
        result = parse_document(self.write("empty.txt", "\ufeff\n"))
        self.assertEqual(result.status, "parsed")
        self.assertIn("empty_file", result.warnings)
        self.assertEqual(result.document.units, [])

    def test_txt_invalid_utf8(self):
        result = parse_document(self.write("bad.txt", b"\xff\xfe\x00", binary=True))
        self.assertEqual((result.status, result.warnings), ("failed", ["encoding_not_utf8"]))

    def test_markdown_headings_preserved(self):
        result = parse_document(self.write("a.md", "# 产品\n\n## 证据\n\n保留来源。"))
        self.assertIn("# 产品", result.document.normalized_text)
        self.assertEqual(result.document.units[-1].heading_path, ["产品", "证据"])

    def test_markdown_table_and_list_preserved(self):
        source = "# 能力\n\n- 证据追踪\n\n| 项目 | 状态 |\n| --- | --- |\n| 来源 | 保留 |"
        result = parse_document(self.write("a.md", source))
        self.assertIn("- 证据追踪", result.document.normalized_text)
        self.assertIn("| --- | --- |", result.document.normalized_text)

    def test_markdown_fenced_heading_is_not_section(self):
        source = "# 正文\n\n```md\n# 代码里的标题\n```\n\n说明"
        result = parse_document(self.write("a.md", source))
        self.assertEqual(len(result.document.units), 1)
        self.assertEqual(result.document.units[0].heading_path, ["正文"])

    def test_docx_paragraph(self):
        path = self.docx(build=lambda d: d.add_paragraph("支持证据追踪"))
        self.assertIn("支持证据追踪", parse_document(path).document.normalized_text)

    def test_docx_heading(self):
        def build(document):
            document.add_heading("产品", level=1)
            document.add_heading("证据", level=2)
            document.add_paragraph("保留来源")
        unit = parse_document(self.docx(build=build)).document.units[0]
        self.assertEqual(unit.heading_path, ["产品", "证据"])
        self.assertTrue(unit.text.startswith("# 产品\n## 证据"))

    def test_docx_table(self):
        def build(document):
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text, table.cell(0, 1).text = "能力", "状态"
            table.cell(1, 0).text, table.cell(1, 1).text = "证据", "支持"
        text = parse_document(self.docx(build=build)).document.normalized_text
        self.assertIn("| 能力 | 状态 |", text)
        self.assertIn("| 证据 | 支持 |", text)

    def test_empty_docx(self):
        result = parse_document(self.docx())
        self.assertEqual(result.status, "parsed")
        self.assertIn("empty_document", result.warnings)

    def test_docx_list(self):
        def build(document):
            document.add_paragraph("保留证据", style="List Bullet")
            document.add_paragraph("生成报告", style="List Number")
        text = parse_document(self.docx(build=build)).document.normalized_text
        self.assertIn("- 保留证据", text)
        self.assertIn("1. 生成报告", text)

    def test_docx_reading_order(self):
        def build(document):
            document.add_paragraph("第一段")
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "中间表格"
            document.add_paragraph("最后一段")
        text = parse_document(self.docx(build=build)).document.normalized_text
        self.assertLess(text.index("第一段"), text.index("中间表格"))
        self.assertLess(text.index("中间表格"), text.index("最后一段"))

    def test_pdf_text(self):
        result = parse_document(FIXTURES / "demo.pdf")
        self.assertEqual(result.status, "parsed")
        self.assertIn("竞争事件证据追踪", result.document.normalized_text)

    def test_pdf_multiple_pages(self):
        result = parse_document(FIXTURES / "demo.pdf")
        self.assertEqual([u.page_number for u in result.document.units], [1, 2])

    def test_pdf_page_labels_in_text(self):
        result = parse_document(FIXTURES / "demo.pdf")
        self.assertTrue(result.document.units[0].text.startswith("# Page 1"))
        self.assertTrue(result.document.units[1].text.startswith("# Page 2"))

    def test_blank_pdf_is_scanned_or_no_text(self):
        path = self.root / "blank.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        with path.open("wb") as stream:
            writer.write(stream)
        result = parse_document(path)
        self.assertEqual(result.status, "scanned_or_no_text")
        self.assertEqual(result.document.units, [])

    def test_corrupt_pdf(self):
        result = parse_document(self.write("bad.pdf", b"not pdf", binary=True))
        self.assertEqual((result.status, result.warnings), ("failed", ["invalid_pdf"]))

    def test_xlsx_single_sheet(self):
        def build(workbook):
            sheet = workbook.active
            sheet.title = "能力"
            sheet.append(["能力", "状态"])
            sheet.append(["证据追踪", "支持"])
        result = parse_document(self.xlsx(build=build))
        self.assertEqual(result.document.units[0].sheet_name, "能力")
        self.assertEqual(result.document.units[0].cell_range, "A1:B2")

    def test_xlsx_multiple_sheets(self):
        result = parse_document(FIXTURES / "demo.xlsx")
        self.assertEqual([u.sheet_name for u in result.document.units], ["产品能力", "额度示例"])

    def test_xlsx_header_and_rows(self):
        text = parse_document(FIXTURES / "demo.xlsx").document.units[0].text
        self.assertIn("| 模块 | 能力 | 状态 |", text)
        self.assertIn("| 事件 | 竞争事件证据追踪 | 已支持 |", text)

    def test_xlsx_empty_sheet_warning(self):
        result = parse_document(self.xlsx())
        self.assertEqual(result.status, "parsed")
        self.assertIn("empty_sheet:Sheet", result.warnings)
        self.assertIn("workbook_has_no_data", result.warnings)

    def test_xlsx_formula_is_not_claimed_as_value_without_cache(self):
        def build(workbook):
            sheet = workbook.active
            sheet.append(["输入", "结果"])
            sheet.append([2, "=A2*2"])
        result = parse_document(self.xlsx(build=build))
        self.assertIn("[公式: =A2*2]", result.document.normalized_text)
        self.assertIn("formula_without_cached_value:Sheet!B2", result.warnings)
        self.assertEqual(result.document.units[0].metadata["formula_cells"][0]["value_source"], "formula_only")

    def test_xlsx_merged_cell(self):
        def build(workbook):
            sheet = workbook.active
            sheet.merge_cells("A1:B1")
            sheet["A1"] = "合并标题"
            sheet.append(["证据", "支持"])
        unit = parse_document(self.xlsx(build=build)).document.units[0]
        self.assertIn("合并标题", unit.text)
        self.assertEqual(unit.metadata["merged_ranges"], ["A1:B1"])

    def test_unsupported_suffix(self):
        result = parse_document(self.write("legacy.doc", "old"))
        self.assertEqual((result.status, result.warnings), ("unsupported", ["unsupported_file_type"]))

    def test_file_not_found(self):
        result = parse_document(self.root / "missing.pdf")
        self.assertEqual((result.status, result.warnings), ("failed", ["file_not_found"]))

    def test_file_size_limit(self):
        result = parse_document(self.write("large.txt", "123456"), max_file_bytes=3)
        self.assertEqual(result.warnings, ["file_size_limit_exceeded"])

    def test_text_length_limit(self):
        result = parse_document(self.write("long.md", "# T\n\n123456"), max_text_chars=5)
        self.assertEqual(result.warnings, ["text_length_limit_exceeded"])

    def test_invalid_docx_container(self):
        result = parse_document(self.write("bad.docx", b"not zip", binary=True))
        self.assertEqual(result.warnings, ["invalid_office_container"])

    def test_invalid_xlsx_container(self):
        result = parse_document(self.write("bad.xlsx", b"not zip", binary=True))
        self.assertEqual(result.warnings, ["invalid_office_container"])

    def test_zip_compression_ratio_limit(self):
        path = self.root / "bomb.docx"
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("huge.xml", b"0" * 5_000_000)
        self.assertEqual(parse_document(path).warnings, ["archive_compression_ratio_exceeded"])

    def test_provenance_converts_to_knowledge_document(self):
        path = FIXTURES / "demo.pdf"
        result = parse_document(path)
        documents = prepare_for_knowledge(result, path, document_id="demo", product_id="pia")
        self.assertEqual([d.metadata["page_number"] for d in documents], [1, 2])
        self.assertTrue(all(d.source_reference == str(path.resolve()) for d in documents))

    def test_nonparsed_result_creates_no_knowledge(self):
        path = self.root / "missing.pdf"
        self.assertEqual(prepare_for_knowledge(parse_document(path), path, document_id="x", product_id="p"), [])

    def test_provenance_rejects_different_source_file(self):
        original = self.write("original.txt", "证据追踪")
        other = self.write("other.txt", "不同内容")
        with self.assertRaisesRegex(ValueError, "来源文件不一致"):
            prepare_for_knowledge(parse_document(original), other, document_id="x", product_id="p")

    def test_parse_and_prepare_is_stable_ui_entry(self):
        prepared = parse_and_prepare(FIXTURES / "demo.docx", document_id="demo", product_id="pia")
        self.assertEqual(prepared.status, "parsed")
        self.assertEqual(prepared.document_metadata["file_type"], "docx")
        self.assertTrue(prepared.ingestion_input)

    def test_prepared_markdown_uses_existing_chunker(self):
        prepared = parse_and_prepare(FIXTURES / "demo.md", document_id="demo", product_id="pia")
        chunks = [chunk for document in prepared.ingestion_input for chunk in chunk_document(document)]
        self.assertTrue(any("竞争事件证据追踪" in chunk.heading_path for chunk in chunks))
        self.assertTrue(any(chunk.table_context for chunk in chunks))

    def test_external_link_is_data_only(self):
        def build(workbook):
            sheet = workbook.active
            sheet["A1"] = "外部链接"
            sheet["A2"] = "https://example.com/do-not-open"
            sheet["A2"].hyperlink = "https://example.com/do-not-open"
        path = self.xlsx(build=build)
        with patch("requests.sessions.Session.request") as request, patch("subprocess.run") as process:
            result = parse_document(path)
        self.assertEqual(result.status, "parsed")
        self.assertIn("https://example.com/do-not-open", result.document.normalized_text)
        request.assert_not_called()
        process.assert_not_called()

    def test_source_metadata_has_no_absolute_path(self):
        metadata = parse_document(FIXTURES / "demo.txt").document.source_metadata
        self.assertEqual(metadata["file_name"], "demo.txt")
        self.assertRegex(metadata["file_sha256"], r"^[0-9a-f]{64}$")
        self.assertNotIn("source_reference", metadata)


if __name__ == "__main__":
    unittest.main()
