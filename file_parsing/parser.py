"""五种用户文件的本地解析实现；不联网、不执行宏、不调用 LLM 或 OCR。"""

from __future__ import annotations

import re
import zipfile
import hashlib
from datetime import date, datetime, time
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

from .schemas import ParsedDocument, ParsedUnit, ParsingResult


SUPPORTED_SUFFIXES = {".txt": "txt", ".md": "md", ".docx": "docx", ".pdf": "pdf", ".xlsx": "xlsx"}
DEFAULT_MAX_FILE_BYTES = 20 * 1024 * 1024
DEFAULT_MAX_TEXT_CHARS = 2_000_000
MAX_ZIP_MEMBERS = 5_000
MAX_ZIP_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ZIP_RATIO = 1_000
MAX_XLSX_AREA_CELLS = 100_000
MIN_PDF_TEXT_CHARS = 20


class ParsingFailure(ValueError):
    """只携带可展示的稳定错误码，不保留解析库或本机路径信息。"""


def _result(status, warning, document=None):
    warnings = list(dict.fromkeys(([warning] if isinstance(warning, str) else warning)))
    return ParsingResult(status=status, document=document, warnings=warnings)


def _normalize(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n").strip()


def file_sha256(path: Path) -> str:
    """流式计算文件指纹，供解析结果与后续 Knowledge provenance 互相核对。"""
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _zip_preflight(path: Path) -> None:
    """读取压缩目录并限制成员数、解压体积和压缩比，避免先完整解压恶意 Office 文件。"""
    if not zipfile.is_zipfile(path):
        raise ParsingFailure("invalid_office_container")
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_MEMBERS:
                raise ParsingFailure("archive_member_limit_exceeded")
            if any(item.flag_bits & 0x1 for item in members):
                raise ParsingFailure("encrypted_office_file")
            expanded = sum(item.file_size for item in members)
            compressed = sum(item.compress_size for item in members)
            if expanded > MAX_ZIP_UNCOMPRESSED_BYTES:
                raise ParsingFailure("archive_expanded_size_limit_exceeded")
            if compressed and expanded / compressed > MAX_ZIP_RATIO:
                raise ParsingFailure("archive_compression_ratio_exceeded")
    except (OSError, zipfile.BadZipFile):
        raise ParsingFailure("invalid_office_container") from None


def _markdown_units(text: str) -> list[ParsedUnit]:
    """只识别围栏代码外的 ATX 标题以形成可追踪章节，正文保持原 Markdown 不变。"""
    lines = text.splitlines(keepends=True)
    headings: list[tuple[int, str]] = []
    units, start, path = [], 0, []
    offset, fence = 0, None

    def append_unit(raw: str, heading_path: list[str]) -> None:
        """过滤只有标题的区段，并补回完整标题链供既有结构切片继续使用。"""
        if heading_path:
            body = re.sub(r"^#{1,6}\s+.+?\s*\n+", "", raw.lstrip(), count=1)
            if not body.strip():
                return
            prefix = "\n".join("#" * (level + 1) + " " + label
                               for level, label in enumerate(heading_path))
            value = prefix + "\n\n" + body.rstrip()
        else:
            value = raw.rstrip()
        if value.strip():
            units.append(ParsedUnit(text=value, heading_path=heading_path))

    for line in lines:
        marker = re.match(r"^\s*(`{3,}|~{3,})", line)
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line.rstrip("\n")) if fence is None else None
        if heading:
            append_unit(text[start:offset], path)
            level, label = len(heading[1]), heading[2]
            headings = [(n, h) for n, h in headings if n < level] + [(level, label)]
            path, start = [h for _, h in headings], offset
        if marker:
            if fence is None:
                fence = marker[1]
            elif marker[1][0] == fence[0] and len(marker[1]) >= len(fence):
                fence = None
        offset += len(line)
    append_unit(text[start:], path)
    return units


def _plain_text(path: Path, kind: str) -> tuple[str, list[ParsedUnit], str, list[str]]:
    try:
        text = _normalize(path.read_text(encoding="utf-8-sig"))
    except UnicodeDecodeError:
        raise ParsingFailure("encoding_not_utf8") from None
    if not text:
        return "", [], path.stem, ["empty_file"]
    if kind == "md":
        units = _markdown_units(text)
        title_match = re.search(r"^#\s+(.+?)\s*$", text, re.MULTILINE)
        return text, units, title_match[1].strip() if title_match else path.stem, []
    return text, [ParsedUnit(text=text)], path.stem, []


def _table_markdown(table: Table) -> tuple[str, int]:
    rows = []
    for row in table.rows:
        values = [cell.text.replace("\r", "").replace("\n", "<br>").replace("|", "\\|").strip()
                  for cell in row.cells]
        if any(values):
            rows.append(values)
    if not rows:
        return "", 0
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    markdown = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    markdown.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return "\n".join(markdown), len(rows)


def _docx(path: Path) -> tuple[str, list[ParsedUnit], str, list[str]]:
    _zip_preflight(path)
    try:
        document = Document(path)
    except Exception:
        raise ParsingFailure("invalid_docx") from None
    warnings = []
    if document.inline_shapes:
        warnings.append("images_ignored_no_ocr")
    if any(p.text.strip() for section in document.sections for p in (*section.header.paragraphs, *section.footer.paragraphs)):
        warnings.append("headers_footers_ignored")
    try:
        if document.element.xpath(".//w:txbxContent"):
            warnings.append("floating_text_ignored")
    except Exception:
        pass
    headings: list[tuple[int, str]] = []
    units, blocks, unit_path = [], [], []

    def flush():
        if not any(block.strip() for block in blocks):
            return
        prefix = "\n".join("#" * (i + 1) + " " + name for i, name in enumerate(unit_path))
        body = "\n\n".join(blocks)
        units.append(ParsedUnit(text=(prefix + "\n\n" + body).strip(), heading_path=unit_path))

    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            value = item.text.strip()
            if not value:
                continue
            style = item.style.name if item.style else ""
            match = re.fullmatch(r"Heading\s+([1-6])", style, re.IGNORECASE)
            if match:
                flush()
                blocks = []
                level = int(match[1])
                headings = [(n, h) for n, h in headings if n < level] + [(level, value)]
                unit_path = [h for _, h in headings]
                continue
            numbered = style.lower().startswith("list number")
            listed = numbered or style.lower().startswith("list") or item._p.pPr is not None and item._p.pPr.numPr is not None
            blocks.append(("1. " if numbered else "- ") + value if listed else value)
        elif isinstance(item, Table):
            table_text, _ = _table_markdown(item)
            if table_text:
                blocks.append(table_text)
    flush()
    title = (document.core_properties.title or "").strip() or (units[0].heading_path[0] if units and units[0].heading_path else path.stem)
    text = "\n\n".join(unit.text for unit in units)
    if not text:
        warnings.append("empty_document")
    return text, units, title, warnings


def _pdf(path: Path) -> tuple[str, list[ParsedUnit], str, list[str], bool]:
    try:
        reader = PdfReader(path, strict=False)
        if reader.is_encrypted and not reader.decrypt(""):
            raise ParsingFailure("encrypted_pdf")
        units, warnings = [], []
        for number, page in enumerate(reader.pages, 1):
            text = _normalize(page.extract_text() or "")
            if not text:
                warnings.append(f"empty_pdf_page:{number}")
                continue
            units.append(ParsedUnit(text=f"# Page {number}\n\n{text}", page_number=number))
    except ParsingFailure:
        raise
    except Exception:
        raise ParsingFailure("invalid_pdf") from None
    visible = sum(len(re.sub(r"\s", "", unit.text.partition("\n\n")[2])) for unit in units)
    scanned = visible < MIN_PDF_TEXT_CHARS
    if scanned:
        warnings.append("scanned_or_no_text")
        units = []
    return "\n\n".join(unit.text for unit in units), units, path.stem, warnings, scanned


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value).replace("\r", "").replace("\n", "<br>").replace("|", "\\|")


def _xlsx(path: Path) -> tuple[str, list[ParsedUnit], str, list[str]]:
    _zip_preflight(path)
    try:
        formulas = load_workbook(path, read_only=False, data_only=False, keep_links=False)
        values = load_workbook(path, read_only=False, data_only=True, keep_links=False)
    except Exception:
        raise ParsingFailure("invalid_xlsx") from None
    warnings, units = [], []
    try:
        for sheet in formulas.worksheets:
            value_sheet = values[sheet.title]
            cells = [cell for cell in sheet._cells.values() if cell.value is not None]
            if not cells:
                warnings.append(f"empty_sheet:{sheet.title}")
                continue
            min_row, max_row = min(c.row for c in cells), max(c.row for c in cells)
            min_col, max_col = min(c.column for c in cells), max(c.column for c in cells)
            if (max_row - min_row + 1) * (max_col - min_col + 1) > MAX_XLSX_AREA_CELLS:
                raise ParsingFailure("xlsx_used_range_limit_exceeded")
            rows, formula_cells = [], []
            for row in range(min_row, max_row + 1):
                values_row = []
                for column in range(min_col, max_col + 1):
                    cell = sheet.cell(row, column)
                    value = cell.value
                    if cell.data_type == "f":
                        cached = value_sheet.cell(row, column).value
                        source = "cached_value" if cached is not None else "formula_only"
                        formula_cells.append({"cell": cell.coordinate, "formula": str(value), "value_source": source})
                        if cached is None:
                            warnings.append(f"formula_without_cached_value:{sheet.title}!{cell.coordinate}")
                            value = f"[公式: {value}]"
                        else:
                            value = cached
                    values_row.append(_cell_text(value))
                rows.append(values_row)
            width = max_col - min_col + 1
            table = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
            table.extend("| " + " | ".join(row) + " |" for row in rows[1:])
            cell_range = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
            merged = [str(item) for item in sheet.merged_cells.ranges]
            unit = ParsedUnit(text=f"# Sheet: {sheet.title}\n\n" + "\n".join(table),
                              heading_path=[f"Sheet: {sheet.title}"], sheet_name=sheet.title,
                              cell_range=cell_range, metadata={"formula_cells": formula_cells,
                                                               "merged_ranges": merged})
            units.append(unit)
    finally:
        formulas.close()
        values.close()
    text = "\n\n".join(unit.text for unit in units)
    if not text:
        warnings.append("workbook_has_no_data")
    return text, units, path.stem, warnings


def parse_document(path: str | Path, *, max_file_bytes: int = DEFAULT_MAX_FILE_BYTES,
                   max_text_chars: int = DEFAULT_MAX_TEXT_CHARS) -> ParsingResult:
    """解析一个明确文件并返回统一结果。

    输入是单个本地路径和资源上限；函数先验证文件/压缩容器，再调用相应本地库，
    输出状态、结构化正文、位置单元和稳定 warning。解析失败不返回底层异常正文。
    """
    if type(max_file_bytes) is not int or max_file_bytes <= 0 or type(max_text_chars) is not int or max_text_chars <= 0:
        raise ValueError("解析上限必须是正整数。")
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        return _result("unsupported", "unsupported_file_type")
    if not file_path.is_file():
        return _result("failed", "file_not_found")
    try:
        size = file_path.stat().st_size
    except OSError:
        return _result("failed", "file_unreadable")
    if size > max_file_bytes:
        return _result("failed", "file_size_limit_exceeded")
    kind = SUPPORTED_SUFFIXES[suffix]
    try:
        scanned = False
        if kind in {"txt", "md"}:
            text, units, title, warnings = _plain_text(file_path, kind)
        elif kind == "docx":
            text, units, title, warnings = _docx(file_path)
        elif kind == "pdf":
            text, units, title, warnings, scanned = _pdf(file_path)
        else:
            text, units, title, warnings = _xlsx(file_path)
        if len(text) > max_text_chars:
            raise ParsingFailure("text_length_limit_exceeded")
        document = ParsedDocument(file_name=file_path.name, file_type=kind, title=title,
                                  normalized_text=text, units=units,
                                  source_metadata={"file_name": file_path.name, "file_type": kind,
                                                   "file_size_bytes": size, "file_sha256": file_sha256(file_path),
                                                   "unit_count": len(units)},
                                  warnings=list(dict.fromkeys(warnings)))
        return _result("scanned_or_no_text" if scanned else "parsed", document.warnings, document)
    except ParsingFailure as exc:
        return _result("failed", str(exc))
    except (OSError, ValueError):
        return _result("failed", "file_parse_failed")
